"""corpus.py — 등급(C/S/O)·포맷·난이도별 라벨된 합성 테스트셋 생성.

설계 원칙
---------
1. **정책 기반 라벨링** — 분류기 규칙이 아니라 N²SF 정책 정의로 등급을 부여한다.
     O(공개): 개인식별자·기밀 라벨 없음.
     S(민감): 제한적 개인정보(이름·연락처·이메일·주소·사업자번호). 강식별자/기밀어 없음.
     C(기밀): 강식별자(주민번호·카드·여권·계좌·API/AWS 키) **또는** 대량 PII(≥10건) **또는**
             기밀 키워드(극비/대외비/기밀/Confidential) 포함.
2. **동일 콘텐츠 → 다중 포맷** — 같은 논리 문서를 txt/csv/json/md/docx/xlsx/pptx/hwpx/pdf 로
   렌더. 포맷별 결과 차이 = 추출/탐지 레이어 버그를 분리 검출.
3. **난이도 분류** — normal(정상) / hard_neg(오탐 유발) / format_stress(추출 스트레스).

오피스/HWPX 는 표준 라이브러리 zipfile 로 분류기가 읽는 XML(`<w:t>`/`<a:t>`/`<hp:t>`)을 직접 생성한다
(분류기의 "외부 의존 0" 추출 철학과 일치). xlsx 는 openpyxl, pdf 는 reportlab 사용.
"""
from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from .synth import Synth

ALL_FORMATS = ["txt", "csv", "json", "md", "docx", "xlsx", "pptx", "hwpx", "pdf"]
GRADES = ["O", "S", "C"]


@dataclass
class LogicalDoc:
    doc_id: str
    grade: str                     # 정책 라벨 O/S/C
    category: str                  # normal | hard_neg | format_stress
    locale: str
    title: str
    paragraphs: List[str] = field(default_factory=list)
    rows: Optional[List[List[str]]] = None     # 표형 데이터(명부 등)
    notes: Optional[str] = None                # pptx 노트 / docx 각주에 숨길 내용
    expected: Dict[str, int] = field(default_factory=dict)   # 주입한 PII 타입별 건수(감사용)


# ════════════════════════════════════════════════════════════════════════
# 1. 논리 문서 생성 (등급·난이도별)
# ════════════════════════════════════════════════════════════════════════
class CorpusGen:
    def __init__(self, seed: int = 0, locale: str = "ko", tag: str = ""):
        self.s = Synth(seed)
        self.r = self.s.r
        self.locale = locale
        self.tag = tag            # 반복 간 doc_id 충돌 방지용 접두(예: 'L2i7')
        self._n = 0

    def _id(self, grade: str, cat: str) -> str:
        self._n += 1
        pre = f"{self.tag}-" if self.tag else ""
        return f"{pre}{grade}-{cat}-{self._n:04d}"

    # ── OPEN ────────────────────────────────────────────────────────────
    def open_doc(self, cat: str = "normal") -> LogicalDoc:
        topic = self.r.choice([
            ("분기 사업 현황 공개 보고", ["당사는 이번 분기 매출이 전년 대비 안정적인 성장세를 보였습니다.",
                                  "신규 오픈소스 라이선스 가이드라인을 사내 위키에 공개했습니다.",
                                  "전 직원 대상 공개 세미나가 다음 주 대강당에서 진행됩니다."]),
            ("공개 채용 공고", ["당사는 백엔드 엔지니어를 공개 채용합니다.",
                          "지원 자격: 관련 경력 3년 이상, 우대사항은 공고 본문을 참고하세요.",
                          "접수는 채용 홈페이지를 통해 진행되며 별도 첨부는 불필요합니다."]),
            ("오픈소스 릴리스 노트", ["v2.3.0 에서 성능이 개선되고 버그가 수정되었습니다.",
                            "본 릴리스는 공개 저장소에서 누구나 내려받을 수 있습니다.",
                            "기여 가이드는 CONTRIBUTING 문서를 확인하세요."]),
        ])
        title, paras = topic
        doc = LogicalDoc(self._id("O", cat), "O", cat, self.locale, title, list(paras))
        if cat == "hard_neg":
            # 오탐 유발: 기밀어를 부정문맥에 / 13자리 제품코드 / 공개 대표번호
            trap = self.r.choice([
                f"본 자료에는 기밀이 포함되어 있지 않으며 자유롭게 배포 가능합니다.",
                f"제품 코드 {self.s.product_code()} 는 공개 카탈로그 식별자입니다.",
                f"주문번호 {self.s.order_no()} 로 배송 상태를 조회할 수 있습니다.",
            ])
            doc.paragraphs.append(trap)
        return doc

    # ── SENSITIVE ───────────────────────────────────────────────────────
    # 정책: 제한적 개인정보(단일 담당자 수준). 강식별자 없음, 대량 PII(≥10) 아님.
    # 다인 명부(연락처 다수)는 정책상 C(대량 PII)로 분류되므로 여기서 생성하지 않는다.
    def sensitive_doc(self, cat: str = "normal") -> LogicalDoc:
        title = self.r.choice(["행사 안내 및 문의처", "공지 — 담당자 연락 안내", "교육 신청 안내"])
        doc = LogicalDoc(self._id("S", cat), "S", cat, self.locale, title)
        exp: Dict[str, int] = {}
        person = self.s.name()
        contact = f"담당자 {person}, 연락처 {self.s.phone()}, 이메일 {self.s.email()}"
        exp.update({"KR_NAME": 1, "KR_PHONE": 1, "EMAIL_ADDRESS": 1})
        body = [f"{title} 안내드립니다. 자세한 사항은 담당자에게 문의 바랍니다.", contact]
        if cat == "format_stress":
            # 연락처를 본문이 아닌 노트(pptx 노트 / docx 각주)로 → 추출 경로 스트레스
            doc.paragraphs = [f"{title} 안내드립니다. 문의는 첨부된 담당자 정보를 참고하세요."]
            doc.notes = contact
        else:
            doc.paragraphs = body
        doc.expected = exp
        return doc

    # ── CONFIDENTIAL ────────────────────────────────────────────────────
    def confidential_doc(self, cat: str = "normal") -> LogicalDoc:
        kind = self.r.choice(["strong_id", "keyword", "bulk", "credentials"])
        exp: Dict[str, int] = {}
        if kind == "bulk":
            title = "전 직원 비상연락망"
            doc = LogicalDoc(self._id("C", cat), "C", cat, self.locale, title)
            n = self.r.randint(10, 16)      # 대량 PII → C escalation
            rows = [["성명", "휴대폰", "이메일"]]
            for _ in range(n):
                rows.append([self.s.name(), self.s.phone(), self.s.email()])
            doc.rows = rows
            doc.paragraphs = [f"{title} ({n}명). 외부 유출 금지."]
            exp.update({"KR_NAME": n, "KR_PHONE": n, "EMAIL_ADDRESS": n})
        elif kind == "keyword":
            kw = self.r.choice(["대외비", "극비", "기밀", "Confidential"])
            title = f"[{kw}] 경영 전략 검토 자료"
            doc = LogicalDoc(self._id("C", cat), "C", cat, self.locale, title)
            doc.paragraphs = [f"{kw}. 본 문서는 지정된 임직원만 열람할 수 있습니다.",
                              f"담당 {self.s.name()} ({self.s.phone()}).",
                              "전략 방향성 및 투자 우선순위는 본문을 참조하십시오."]
            exp.update({"KR_NAME": 1, "KR_PHONE": 1})
        elif kind == "credentials":
            title = "운영 서버 접속 정보"
            doc = LogicalDoc(self._id("C", cat), "C", cat, self.locale, title)
            doc.paragraphs = [f"운영계 접속 키: {self.s.api_key()}",
                              f"백업 자격증명: {self.s.aws_key()}",
                              f"담당 {self.s.name()} {self.s.email()}."]
            exp.update({"GENERIC_API_KEY": 1, "AWS_ACCESS_KEY": 1, "KR_NAME": 1, "EMAIL_ADDRESS": 1})
        else:  # strong_id
            title = self.r.choice(["임직원 신원 확인 자료", "급여 계좌 등록 명부", "해외 출장자 여권 정보"])
            doc = LogicalDoc(self._id("C", cat), "C", cat, self.locale, title)
            strong = self.r.choice(["rrn", "card", "passport", "account"])
            person = self.s.name()
            if strong == "rrn":
                line = f"성명 {person} 주민등록번호 {self.s.rrn()}"
                exp["KR_RRN"] = 1
            elif strong == "card":
                line = f"성명 {person} 신용카드번호 {self.s.credit_card()}"
                exp["CREDIT_CARD"] = 1
            elif strong == "passport":
                line = f"성명 {person} 여권번호 {self.s.passport()}"
                exp["KR_PASSPORT"] = 1
            else:
                line = f"예금주 {person} 계좌 {self.s.account()}"
                exp["KR_ACCOUNT"] = 1
            exp["KR_NAME"] = exp.get("KR_NAME", 0) + 1
            doc.paragraphs = [f"{title}. 개인정보보호법에 따라 비식별 처리 대상입니다.", line,
                              f"연락처 {self.s.phone()}."]
            exp["KR_PHONE"] = 1
        doc.expected = exp
        if cat == "format_stress":
            # 강식별자를 본문이 아닌 표 셀(숫자 셀 가능) / 노트로 이동
            if doc.rows is None and "KR_ACCOUNT" not in exp:
                acct = self.s.account().split()[-1].replace("-", "")    # 패킹된 계좌(숫자셀 후보)
                doc.rows = [["항목", "값"], ["계좌(패킹)", acct]]
                exp["KR_ACCOUNT"] = exp.get("KR_ACCOUNT", 0) + 1
            doc.notes = f"비고: 추가 식별자 주민번호 {self.s.rrn()}"
            exp["KR_RRN"] = exp.get("KR_RRN", 0) + 1
        return doc

    # ── 경계/적대 케이스 (난이도 L2+) ───────────────────────────────────
    def boundary_doc(self, cat: str = "boundary") -> LogicalDoc:
        """연락처 k명 명부 — 정책: 개인식별자 합계 ≥10 이면 C, 미만 S. bulk 임계 자극."""
        k = self.r.randint(4, 8)
        ids = 2 * k                          # phone + email per person
        grade = "C" if ids >= 10 else "S"    # k>=5 → C, k==4 → S
        doc = LogicalDoc(self._id(grade, cat), grade, cat, self.locale, "담당자 연락처 목록")
        rows = [["성명", "연락처", "이메일"]]
        for _ in range(k):
            rows.append([self.s.name(), self.s.phone(), self.s.email()])
        doc.rows = rows
        doc.paragraphs = [f"담당자 연락처 목록 ({k}명)."]
        doc.expected = {"KR_NAME": k, "KR_PHONE": k, "EMAIL_ADDRESS": k}
        return doc

    def open_adversarial_doc(self, cat: str = "adversarial") -> LogicalDoc:
        """공개 문서지만 PII 처럼 보이는 미끼(체크섬 깨진 주민번호·Luhn 카드·긴숫자) 포함 → 정밀도 트랩."""
        doc = LogicalDoc(self._id("O", cat), "O", cat, self.locale, "공개 카탈로그 / 정산 안내")
        trap = self.r.choice([
            f"정산 일련번호 {self.s.near_miss_rrn().replace('-', '')} 로 조회 가능합니다(개인정보 아님).",
            f"상품 코드 {self.s.decoy_card_number()} 는 공개 SKU 식별자입니다.",
            f"문서 일련번호 {self.s.product_code()} / 접수번호 {self.r.randint(10**11, 10**12)}.",
        ])
        doc.paragraphs = ["본 안내는 공개 자료입니다.", trap] + self.s.filler(2)
        return doc

    def confidential_obfuscated_doc(self, cat: str = "obfuscated") -> LogicalDoc:
        """기밀이지만 키워드 난독화(공백/전각) + 미묘한 강식별자 → 재현율 트랩(규칙 회피, 의미는 기밀)."""
        doc = LogicalDoc(self._id("C", cat), "C", cat, self.locale, "내부 검토 자료")
        kw = self.r.choice(["대외비", "극비", "기밀"])
        doc.paragraphs = [f"본 자료는 {self.s.spaced_keyword(kw)} 입니다. 무단 배포를 금합니다.",
                          f"검토 담당 {self.s.name()}.",
                          f"참고 식별 정보: {self.s.fullwidth_digits(self.s.rrn())}"] + self.s.filler(1)
        doc.expected = {"KR_NAME": 1}
        return doc

    # ── 코퍼스 빌드 (난이도 1~4) ────────────────────────────────────────
    def _noisify(self, doc: LogicalDoc, level: int):
        """난이도에 따라 필러 문장을 섞어 신호 희석(추출/탐지 난도↑)."""
        if level >= 3:
            doc.paragraphs = self.s.filler(level - 1) + doc.paragraphs + self.s.filler(level - 1)

    def build(self, per_cell: int = 6, difficulty: int = 1) -> List[LogicalDoc]:
        """등급×난이도 격자로 문서 생성. difficulty 1~4 로 어려운 셀을 점증 추가."""
        makers = {"O": self.open_doc, "S": self.sensitive_doc, "C": self.confidential_doc}
        plan = [("O", "normal"), ("O", "hard_neg"),
                ("S", "normal"), ("S", "format_stress"),
                ("C", "normal"), ("C", "format_stress")]
        special = []   # (callable, cat)
        if difficulty >= 2:
            special += [(self.boundary_doc, "boundary"),
                        (self.open_adversarial_doc, "adversarial")]
        if difficulty >= 3:
            special += [(self.confidential_obfuscated_doc, "obfuscated"),
                        (self.open_adversarial_doc, "adversarial")]
        if difficulty >= 4:
            # format_stress(강식별자를 노트/숫자셀로) 비중 확대 + 난독 기밀 추가
            plan += [("C", "format_stress"), ("S", "format_stress")]
            special += [(self.confidential_obfuscated_doc, "obfuscated")]

        docs: List[LogicalDoc] = []
        for grade, cat in plan:
            for _ in range(per_cell):
                d = makers[grade](cat)
                self._noisify(d, difficulty)
                docs.append(d)
        for maker, _cat in special:
            for _ in range(per_cell):
                d = maker()
                self._noisify(d, difficulty)
                docs.append(d)
        return docs


# ════════════════════════════════════════════════════════════════════════
# 2. 렌더러 (논리 문서 → 포맷별 파일)
# ════════════════════════════════════════════════════════════════════════
def _flat_text(doc: LogicalDoc, include_notes: bool = True) -> str:
    parts = [doc.title, ""]
    parts += doc.paragraphs
    if doc.rows:
        for row in doc.rows:
            parts.append(" ".join(str(c) for c in row))
    if include_notes and doc.notes:
        parts.append("[비고] " + doc.notes)
    return "\n".join(parts)


def _zip_write(path: Path, files: Dict[str, str]):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


def _xml_escape(s: str) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;"))


def render_txt(doc, path):  path.write_text(_flat_text(doc), encoding="utf-8")
def render_md(doc, path):
    body = [f"# {doc.title}", ""]
    body += doc.paragraphs
    if doc.rows:
        body.append("")
        for i, row in enumerate(doc.rows):
            body.append("| " + " | ".join(str(c) for c in row) + " |")
            if i == 0:
                body.append("|" + "---|" * len(row))
    if doc.notes:
        body += ["", f"> 비고: {doc.notes}"]
    path.write_text("\n".join(body), encoding="utf-8")


def render_csv(doc, path):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([doc.title])
    for p in doc.paragraphs:
        w.writerow([p])
    if doc.rows:
        for row in doc.rows:
            w.writerow(row)
    if doc.notes:
        w.writerow(["비고", doc.notes])
    path.write_text(buf.getvalue(), encoding="utf-8")


def render_json(doc, path):
    obj = {"title": doc.title, "paragraphs": doc.paragraphs,
           "rows": doc.rows, "notes": doc.notes}
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def render_docx(doc, path):
    """본문 → word/document.xml, 노트 → word/footnotes.xml(각주, 추출 스트레스)."""
    def para(t):
        return f'<w:p><w:r><w:t xml:space="preserve">{_xml_escape(t)}</w:t></w:r></w:p>'
    body = [para(doc.title)] + [para(p) for p in doc.paragraphs]
    if doc.rows:
        for row in doc.rows:
            body.append(para(" ".join(str(c) for c in row)))
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:body>{"".join(body)}</w:body></w:document>')
    files = {
        "[Content_Types].xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '</Types>',
        "_rels/.rels":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>',
        "word/document.xml": document,
    }
    if doc.notes:
        files["word/footnotes.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:footnote w:id="1"><w:p><w:r><w:t xml:space="preserve">{_xml_escape(doc.notes)}</w:t></w:r></w:p></w:footnote>'
            '</w:footnotes>')
    _zip_write(path, files)


def render_xlsx(doc, path):
    """openpyxl 로 생성. 순수 10~16자리 숫자는 **숫자 셀**로 저장 → 숫자 PII 복원 스트레스."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append([doc.title])
    for p in doc.paragraphs:
        ws.append([p])
    if doc.rows:
        for row in doc.rows:
            out = []
            for c in row:
                cs = str(c)
                if cs.isdigit() and 10 <= len(cs) <= 16:
                    out.append(int(cs))         # 숫자 셀로 저장
                else:
                    out.append(cs)
            ws.append(out)
    if doc.notes:
        ws.append(["비고", doc.notes])
    wb.save(path)


def render_pptx(doc, path):
    """본문 → 슬라이드(<a:t>), 노트 → notesSlide(<a:t>) (발표자 노트 추출 스트레스)."""
    def atext(t):
        return f'<a:p><a:r><a:t>{_xml_escape(t)}</a:t></a:r></a:p>'
    slide_body = [atext(doc.title)] + [atext(p) for p in doc.paragraphs]
    if doc.rows:
        for row in doc.rows:
            slide_body.append(atext(" ".join(str(c) for c in row)))
    slide = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<p:cSld><p:spTree><p:sp><p:txBody>{"".join(slide_body)}</p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
    files = {
        "[Content_Types].xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
            '<Override PartName="/ppt/notesSlides/notesSlide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml"/>'
            '</Types>',
        "_rels/.rels":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>'
            '</Relationships>',
        "ppt/presentation.xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
        "ppt/slides/slide1.xml": slide,
    }
    if doc.notes:
        files["ppt/notesSlides/notesSlide1.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<p:cSld><p:spTree><p:sp><p:txBody>{atext(doc.notes)}</p:txBody></p:sp></p:spTree></p:cSld></p:notes>')
    _zip_write(path, files)


def render_hwpx(doc, path):
    """HWPX: Contents/section0.xml (<hp:t>) + header. 한글 문서 추출 경로 검증."""
    def hp(t):
        return f'<hp:p><hp:run><hp:t>{_xml_escape(t)}</hp:t></hp:run></hp:p>'
    body = [hp(doc.title)] + [hp(p) for p in doc.paragraphs]
    if doc.rows:
        for row in doc.rows:
            body.append(hp(" ".join(str(c) for c in row)))
    if doc.notes:
        body.append(hp("비고 " + doc.notes))
    section = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<hp:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
               f'{"".join(body)}</hp:sec>')
    files = {
        "mimetype": "application/hwp+zip",
        "Contents/header.xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"/>',
        "Contents/section0.xml": section,
    }
    _zip_write(path, files)


def render_pdf(doc, path):
    """reportlab 로 텍스트 PDF 생성(분류기는 pypdf/pdfminer 로 추출)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
        font = "HYSMyeongJo-Medium"
    except Exception:
        font = "Helvetica"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont(font, 11)
    y = 800
    lines = _flat_text(doc).split("\n")
    for line in lines:
        c.drawString(40, y, line[:120])
        y -= 16
        if y < 40:
            c.showPage(); c.setFont(font, 11); y = 800
    c.save()


RENDERERS = {
    "txt": render_txt, "md": render_md, "csv": render_csv, "json": render_json,
    "docx": render_docx, "xlsx": render_xlsx, "pptx": render_pptx,
    "hwpx": render_hwpx, "pdf": render_pdf,
}


def render_all(docs: List[LogicalDoc], out_dir: Path, formats: List[str] = None) -> List[dict]:
    """논리 문서들을 지정 포맷으로 렌더. 코퍼스 매니페스트(파일 행 리스트) 반환."""
    formats = formats or ALL_FORMATS
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for doc in docs:
        for fmt in formats:
            fname = f"{doc.doc_id}.{fmt}"
            fpath = out_dir / fname
            try:
                RENDERERS[fmt](doc, fpath)
            except Exception as exc:        # 렌더 실패는 기록만 (예: 폰트 누락)
                manifest.append({"doc_id": doc.doc_id, "grade": doc.grade,
                                 "category": doc.category, "locale": doc.locale,
                                 "fmt": fmt, "path": str(fpath), "render_error": str(exc),
                                 "expected": doc.expected})
                continue
            manifest.append({"doc_id": doc.doc_id, "grade": doc.grade,
                             "category": doc.category, "locale": doc.locale,
                             "fmt": fmt, "path": str(fpath), "render_error": None,
                             "expected": doc.expected})
    return manifest
